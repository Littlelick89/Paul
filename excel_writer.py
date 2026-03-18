"""Write extracted training data records to an Excel workbook."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

import config


def _col_letter_to_index(letter: str) -> int:
    """Convert a column letter ('A', 'AA', …) to a 1-based integer."""
    return column_index_from_string(letter.upper())


def _find_next_empty_row(ws, start_row: int, key_col: str) -> int:
    """Return the first row >= *start_row* where *key_col* is empty."""
    col_idx = _col_letter_to_index(key_col)
    row = start_row
    while ws.cell(row=row, column=col_idx).value not in (None, ""):
        row += 1
    return row


def _write_headers(ws) -> None:
    """Write column header labels to row 1 of *ws*."""
    col_map = config.EXCEL_COLUMN_MAP
    headers = config.EXCEL_HEADERS
    for field, col_letter in col_map.items():
        label = headers.get(field, field)
        ws.cell(row=1, column=_col_letter_to_index(col_letter), value=label)


def _sheet_has_headers(ws) -> bool:
    """Return True if the first cell of the sheet already has content."""
    first_col = next(iter(config.EXCEL_COLUMN_MAP.values()))
    return ws.cell(row=1, column=_col_letter_to_index(first_col)).value is not None


def write_records(
    records: list[dict[str, Any]],
    excel_path: str | Path,
    sheet_name: str | None = None,
) -> int:
    """Append *records* to *excel_path* and return the number of rows written.

    - Creates the workbook / sheet if they do not exist.
    - Writes headers to row 1 if the sheet is empty.
    - Skips records flagged with ``_skip`` or ``_raw_response``.
    """
    excel_path = Path(excel_path)

    if excel_path.exists():
        wb = load_workbook(excel_path)
    else:
        wb = openpyxl.Workbook()

    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    elif sheet_name:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb.active

    # Write headers if the sheet is brand-new / empty
    if not _sheet_has_headers(ws):
        _write_headers(ws)

    col_map = config.EXCEL_COLUMN_MAP
    # Determine first key column for "next empty row" detection
    first_col = next(iter(col_map.values()))

    rows_written = 0
    for record in records:
        if record.get("_skip"):
            continue
        if "_raw_response" in record:
            print(
                f"[SKIP] 파싱 실패 (페이지 {record.get('_source_page', '?')}): "
                f"{record['_raw_response'][:120]}"
            )
            continue

        next_row = _find_next_empty_row(ws, config.EXCEL_START_ROW, first_col)

        for field, col_letter in col_map.items():
            value = record.get(field)
            ws.cell(
                row=next_row,
                column=_col_letter_to_index(col_letter),
                value=value,
            )

        rows_written += 1

    wb.save(excel_path)
    return rows_written
